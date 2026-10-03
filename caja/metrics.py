"""
Motor del Panel de caja: flujo de fondos REAL del negocio, basado en el
libro diario (`MovimientoCaja`), a diferencia de `budgets.metrics` (que mide
devengado).

Reusa `PERIODS`/`period_bounds` de `budgets.metrics` para no duplicar la
lógica de fechas (semana/mes/año) — `caja` puede depender de `budgets`.
"""

from datetime import timedelta
from decimal import Decimal

from django.urls import reverse
from django.utils import timezone
from django.utils.translation import gettext

from budgets.metrics import PERIODS, period_bounds
from budgets.pdf import format_money

from .models import CuentaCaja, MovimientoCaja

ZERO = Decimal("0")


def build_caja(period: str, now=None) -> dict:
    """
    Devuelve un dict con los números crudos (Decimals) del Panel de caja:
    ingresos/egresos/saldo REALES (movimientos CONFIRMADOS) del período
    [start, end), más el saldo actual por cuenta (foto a HOY) y cuántos
    movimientos quedaron en Borrador.
    """
    if period not in PERIODS:
        period = "month"
    now = now or timezone.localtime(timezone.now())

    start, end = period_bounds(period, now)
    start_date = timezone.localtime(start).date()
    end_date = timezone.localtime(end).date()

    movimientos = list(
        MovimientoCaja.objects.filter(
            estado=MovimientoCaja.Estado.CONFIRMADO,
            fecha__gte=start_date,
            fecha__lt=end_date,
        )
        .select_related("cuenta")
        .order_by("-fecha", "-id")
    )
    ingresos_reales = sum(
        (m.monto for m in movimientos if m.tipo == MovimientoCaja.Tipo.INGRESO), ZERO
    )
    egresos_reales = sum(
        (m.monto for m in movimientos if m.tipo == MovimientoCaja.Tipo.EGRESO), ZERO
    )
    n_ingresos = sum(1 for m in movimientos if m.tipo == MovimientoCaja.Tipo.INGRESO)
    n_egresos = sum(1 for m in movimientos if m.tipo == MovimientoCaja.Tipo.EGRESO)

    cuentas = list(CuentaCaja.objects.filter(is_active=True))
    saldos_por_cuenta = [
        {"nombre": c.nombre, "saldo": c.saldo_actual} for c in cuentas
    ]
    saldo_total = sum((c["saldo"] for c in saldos_por_cuenta), ZERO)

    n_pendientes = MovimientoCaja.objects.filter(
        estado=MovimientoCaja.Estado.BORRADOR
    ).count()

    return {
        "period": period,
        "period_label": PERIODS[period]["label"],
        "start": start,
        "end": end,
        "range_str": (
            f"{timezone.localtime(start).strftime('%d/%m/%Y')} – "
            f"{timezone.localtime(end - timedelta(seconds=1)).strftime('%d/%m/%Y')}"
        ),
        "ingresos_reales": ingresos_reales,
        "n_ingresos": n_ingresos,
        "egresos_reales": egresos_reales,
        "n_egresos": n_egresos,
        "saldo_periodo": ingresos_reales - egresos_reales,
        "saldos_por_cuenta": saldos_por_cuenta,
        "saldo_total": saldo_total,
        "n_pendientes": n_pendientes,
        "movimientos": movimientos,
    }


def _money(value) -> str:
    return "$ " + format_money(value or 0)


def caja_template_context(data: dict) -> dict:
    """Arma el contexto final (ya formateado) para el template del Panel de caja."""
    pendientes_url = reverse("admin:caja_movimientocaja_changelist")
    return {
        "period": data["period"],
        "period_label": data["period_label"],
        "range_str": data["range_str"],
        "ingresos_reales": _money(data["ingresos_reales"]),
        "n_ingresos": data["n_ingresos"],
        "egresos_reales": _money(data["egresos_reales"]),
        "n_egresos": data["n_egresos"],
        "saldo_periodo": _money(data["saldo_periodo"]),
        "saldo_positivo": data["saldo_periodo"] >= 0,
        "saldos_por_cuenta": [
            {"nombre": c["nombre"], "saldo": _money(c["saldo"])}
            for c in data["saldos_por_cuenta"]
        ],
        "saldo_total": _money(data["saldo_total"]),
        "n_pendientes": data["n_pendientes"],
        "pendientes_url": f"{pendientes_url}?estado={MovimientoCaja.Estado.BORRADOR}",
        "movimientos": [
            {
                "fecha": m.fecha.strftime("%d/%m/%Y"),
                "cuenta": m.cuenta.nombre,
                "tipo": m.get_tipo_display(),
                "es_ingreso": m.tipo == MovimientoCaja.Tipo.INGRESO,
                "concepto": m.concepto,
                "monto": _money(m.monto),
            }
            for m in data["movimientos"]
        ],
    }


def export_caja_xlsx(data: dict):
    """Arma el .xlsx (una sola hoja) del Panel de caja y devuelve (filename, bytes)."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="111111")
    head_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True, size=14)

    ws = wb.active
    ws.title = gettext("Caja")
    ws["A1"] = gettext("Panel de caja 3darg — %(period)s (%(range)s)") % {
        "period": str(data["period_label"]),
        "range": data["range_str"],
    }
    ws["A1"].font = title_font

    ws.append([])
    ws.append([gettext("Concepto"), gettext("Monto")])
    style_row = ws.max_row
    for c in range(1, 3):
        cell = ws.cell(row=style_row, column=c)
        cell.fill = head_fill
        cell.font = head_font

    rows = [
        (gettext("Ingresos reales del período"), float(data["ingresos_reales"])),
        (gettext("Egresos reales del período"), float(data["egresos_reales"])),
        (gettext("Saldo del período"), float(data["saldo_periodo"])),
    ]
    for label, value in rows:
        ws.append([label, value])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)

    ws.append([])
    ws.append([gettext("Saldo por cuenta (hoy)"), ""])
    style_row = ws.max_row
    for c in range(1, 3):
        cell = ws.cell(row=style_row, column=c)
        cell.fill = head_fill
        cell.font = head_font
    for cuenta in data["saldos_por_cuenta"]:
        ws.append([cuenta["nombre"], float(cuenta["saldo"])])
    ws.append([gettext("Total"), float(data["saldo_total"])])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)

    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 3, 50)

    buffer = BytesIO()
    wb.save(buffer)
    fname = f"panel_caja_{data['period']}_{timezone.localdate().isoformat()}.xlsx"
    return fname, buffer.getvalue()
